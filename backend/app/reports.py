from collections import defaultdict
import csv
from datetime import datetime
from io import BytesIO, StringIO
from urllib.parse import quote

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import Response, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.calculations.amounts import sfh_amount_from_record, sfh_is_inr_currency, whole_number
from app.dashboard import (
    _direct_amount,
    _direct_sales_channel,
    _amazon_is_cancelled,
    _dsg_is_completed,
    _for_period,
    _load_channel_metrics,
    _month,
    _number,
    _unique_order_count,
    _row_value,
)
from app.database.database import SessionLocal
from app.database.models import AmazonDatasetRow, DirectSalesDatasetRow, DSGDatasetRow, SFHDatasetRow, UploadHistory

router = APIRouter(prefix="/api/reports", tags=["reports"])

CHANNELS = (
    ("DSG", DSGDatasetRow),
    ("SFH", SFHDatasetRow),
    ("Amazon", AmazonDatasetRow),
    ("Direct Sales", DirectSalesDatasetRow),
)
REPORT_TYPES = {"summary", "channel", "category", "product"}
HEADER_FILL = PatternFill("solid", fgColor="17213A")
TITLE_FILL = PatternFill("solid", fgColor="5B5CE2")
# Excel treats ordinary grouping commas as western thousands separators.
# Escaped commas are literal lakh/crore separators; the conditional format is
# used for formula cells whose filtered result can change magnitude.
INDIAN_WHOLE_NUMBER_FORMAT = (
    r'[>=10000000]##\,##\,##\,##0;[>=100000]##\,##\,##0;##,##0'
)
INDIAN_CRORE_FORMAT = r'##\,##\,##\,##0;[Red]-##\,##\,##\,##0;-'
INDIAN_LAKH_FORMAT = r'##\,##\,##0;[Red]-##\,##\,##0;-'
INDIAN_THOUSAND_FORMAT = r'##,##0;[Red]-##,##0;-'
INDIAN_STATE_NAMES = {
    "AN": "Andaman and Nicobar Islands", "AP": "Andhra Pradesh",
    "AR": "Arunachal Pradesh", "AS": "Assam", "BR": "Bihar",
    "CH": "Chandigarh", "CG": "Chhattisgarh", "CT": "Chhattisgarh",
    "DN": "Dadra and Nagar Haveli and Daman and Diu",
    "DD": "Dadra and Nagar Haveli and Daman and Diu", "DL": "Delhi",
    "GA": "Goa", "GJ": "Gujarat", "HR": "Haryana",
    "HP": "Himachal Pradesh", "JK": "Jammu and Kashmir",
    "JH": "Jharkhand", "KA": "Karnataka", "KL": "Kerala",
    "LA": "Ladakh", "LD": "Lakshadweep", "MP": "Madhya Pradesh",
    "MH": "Maharashtra", "MN": "Manipur", "ML": "Meghalaya",
    "MZ": "Mizoram", "NL": "Nagaland", "OD": "Odisha", "OR": "Odisha",
    "PY": "Puducherry", "PB": "Punjab", "RJ": "Rajasthan",
    "SK": "Sikkim", "TN": "Tamil Nadu", "TS": "Telangana",
    "TG": "Telangana", "TR": "Tripura", "UP": "Uttar Pradesh",
    "UK": "Uttarakhand", "UT": "Uttarakhand", "WB": "West Bengal",
}
DIRECT_OVERVIEW_TYPES = {
    "Bulk Sales": "Bulk",
    "Direct Sales": "Retail",
    "Stall Sales": "Stall",
    "Language Lab": "Language Lab",
}
DIRECT_OVERVIEW_TYPE_ORDER = ("Bulk", "Retail", "Stall", "Language Lab")


def _channel_display_name(channel: str) -> str:
    return {
        "DSG": "DSG",
        "SFH": "SFH",
    }.get(channel, channel)


def _set_indian_whole_number_format(cell) -> None:
    """Keep the cell numeric while displaying Indian lakh/crore grouping."""
    if cell.data_type == "f" or not isinstance(cell.value, (int, float)):
        cell.number_format = INDIAN_WHOLE_NUMBER_FORMAT
        return
    magnitude = abs(cell.value)
    cell.number_format = (
        INDIAN_CRORE_FORMAT
        if magnitude >= 10_000_000
        else INDIAN_LAKH_FORMAT
        if magnitude >= 100_000
        else INDIAN_THOUSAND_FORMAT
    )


def _matches(month_key: str, grain: str, period: str, year: int) -> bool:
    row_year, row_month = (int(value) for value in month_key.split("-"))
    if grain == "monthly":
        selected_months = {int(value) for value in period.split(",")}
        return row_year == year and row_month in selected_months
    if grain == "quarterly":
        return row_year == year and ((row_month - 1) // 3) + 1 == int(period)
    return row_year == year


def _period_label(grain: str, period: str, year: int) -> str:
    if grain == "monthly":
        selected = sorted({int(value) for value in period.split(",")})
        labels = [datetime(year, value, 1).strftime("%b") for value in selected]
        if len(labels) == 1:
            month_label = labels[0]
        elif len(selected) <= 3:
            month_label = ", ".join(labels)
        elif selected == list(range(selected[0], selected[-1] + 1)):
            month_label = f"{labels[0]} - {labels[-1]}"
        else:
            month_label = f"{len(selected)} Selected Months"
        return f"{month_label} - {str(year)[-2:]}"
    return str(year)


def _filtered_rows(database, grain: str, period: str, year: int):
    upload_dates = {row.upload_id: row.uploaded_at for row in database.query(UploadHistory).all()}
    fallback = datetime.now()
    result = []
    for channel, model in CHANNELS:
        for row in database.query(model).all():
            if channel == "DSG" and not _dsg_is_completed(row):
                continue
            if channel == "Amazon" and _amazon_is_cancelled(row):
                continue
            month_key = _month(row.row_data, upload_dates.get(row.upload_id, fallback))
            if _matches(month_key, grain, period, year):
                result.append((channel, row, month_key))
    return result


def _amount(channel: str, row) -> float:
    if channel == "SFH":
        return _number(sfh_amount_from_record(row.row_data))
    if channel == "Direct Sales":
        return _direct_amount(row)
    return _number(row.amount)


def _invoice_amount(row) -> float:
    value = _row_value(row.row_data, (
        "invoice amount", "total invoice amount", "order total amount", "order total",
        "actual amount", "total", "grand total", "amount",
    ))
    return _number(value if value is not None else getattr(row, "amount", None))


def _quantity(channel: str, row) -> float:
    value = _row_value(
        row.row_data,
        ("quantity", "qty", "category quantity", "mapped quantity"),
    )
    if value is not None and str(value).strip() != "":
        return _number(value)
    # SFH exports are transaction-level: one row is one purchased course.
    # Current SFH source files have Course/Description but no Quantity field.
    if channel in {"SFH", "Amazon"} and (row.product_name or getattr(row, "course", None)):
        return 1.0
    return 0.0


def _indian_integer(value: int) -> str:
    sign = "-" if value < 0 else ""
    digits = str(abs(value))
    if len(digits) <= 3:
        return sign + digits
    tail = digits[-3:]
    head = digits[:-3]
    pairs = []
    while head:
        pairs.append(head[-2:])
        head = head[:-2]
    return sign + ",".join([*reversed(pairs), tail])


def _direct_overview_source(database):
    upload_dates = {
        item.upload_id: item.uploaded_at
        for item in database.query(UploadHistory).all()
    }
    fallback = datetime.now()
    source = []
    seen_ids = set()
    for row in database.query(DirectSalesDatasetRow).all():
        if row.id in seen_ids:
            raise HTTPException(status_code=500, detail="Duplicate Direct Sales database row detected.")
        seen_ids.add(row.id)
        month_key = _month(row.row_data, upload_dates.get(row.upload_id, fallback))
        row_year, row_month = (int(value) for value in month_key.split("-"))
        source.append({
            "id": row.id,
            "year": row_year,
            "month": row_month,
            "type": DIRECT_OVERVIEW_TYPES[_direct_sales_channel(row)],
            "product": str(row.product_name or "Unmapped").strip() or "Unmapped",
            "order": _order_identifier("Direct Sales", row),
            "quantity": _quantity("Direct Sales", row),
            "invoice_quantity": _number(
                _row_value(row.row_data, ("mapped quantity",))
                or _quantity("Direct Sales", row)
            ),
            "amount": _direct_amount(row),
        })
    return source


def _build_direct_sales_overview(
    database,
    *,
    years: set[int],
    months: set[int],
    types: set[str],
    products: set[str],
) -> dict[str, object]:
    valid_types = set(DIRECT_OVERVIEW_TYPES.values())
    if not years or not months or not types:
        raise HTTPException(status_code=422, detail="Select at least one year, month, and type.")
    if not types <= valid_types or any(month < 1 or month > 12 for month in months):
        raise HTTPException(status_code=422, detail="Invalid Direct Sales Overview filters.")

    filtered = [
        item for item in _direct_overview_source(database)
        if item["year"] in years
        and item["month"] in months
        and item["type"] in types
        and (not products or item["product"] in products)
    ]
    grouped = defaultdict(lambda: {"orders": set(), "invoice_quantities": {}, "quantity": 0.0, "amount": 0.0})
    for item in filtered:
        key = (item["year"], item["month"], item["type"], item["product"])
        if item["order"]:
            grouped[key]["orders"].add(item["order"])
        invoice_key = item["order"] or f"row:{item['id']}"
        grouped[key]["invoice_quantities"][invoice_key] = item["invoice_quantity"]
        grouped[key]["quantity"] += item["quantity"]
        grouped[key]["amount"] += item["amount"]

    ordered = sorted(grouped.items(), key=lambda item: (item[0][0], item[0][1], item[0][2], item[0][3].casefold()))
    display_quantities = [0] * len(ordered)
    display_amounts = [0] * len(ordered)
    indexes_by_type = defaultdict(list)
    for index, ((_, _, type_name, _), _) in enumerate(ordered):
        indexes_by_type[type_name].append(index)
    type_order = DIRECT_OVERVIEW_TYPE_ORDER
    raw_type_totals = [
        sum(item["amount"] for item in filtered if item["type"] == type_name)
        for type_name in type_order
    ]
    reconciled_type_totals = dict(zip(
        type_order,
        _reconciled_whole_numbers(raw_type_totals),
    ))
    for type_name, indexes in indexes_by_type.items():
        quantities = _reconciled_whole_numbers([ordered[index][1]["quantity"] for index in indexes])
        amounts = _reconciled_whole_numbers(
            [ordered[index][1]["amount"] for index in indexes],
            target=reconciled_type_totals[type_name],
        )
        for position, index in enumerate(indexes):
            display_quantities[index] = quantities[position]
            display_amounts[index] = amounts[position]

    rows = []
    for index, (((year, month, type_name, product), values)) in enumerate(ordered):
        rows.append({
            "year": year,
            "month": datetime(year, month, 1).strftime("%b"),
            "type": type_name,
            "product": product,
            "orders": len(values["orders"]),
            "invoice_quantity": whole_number(sum(values["invoice_quantities"].values())),
            "quantity": display_quantities[index],
            "without_tax_total": display_amounts[index],
        })
    keys = {(row["year"], row["month"], row["type"], row["product"]) for row in rows}
    if len(keys) != len(rows):
        raise HTTPException(status_code=500, detail="Duplicate rows detected in Direct Sales Overview.")
    totals = {
        type_name: sum(row["without_tax_total"] for row in rows if row["type"] == type_name)
        for type_name in DIRECT_OVERVIEW_TYPE_ORDER
    }
    expected = {type_name: reconciled_type_totals[type_name] for type_name in type_order}
    if totals != expected:
        raise HTTPException(status_code=500, detail="Direct Sales Overview failed dashboard reconciliation.")
    return {
        "rows": rows,
        "row_count": len(rows),
        "source_records": len(filtered),
        "totals": totals,
        "overall_total": sum(totals.values()),
        "validated": True,
    }


def _location(channel: str, data: dict) -> str:
    if channel == "DSG":
        value = _row_value(data, ("state code (billing)",))
        code = str(value or "").strip().upper()
        return INDIAN_STATE_NAMES.get(code, str(value or "").strip())
    if channel == "SFH":
        return str(_row_value(data, ("place of supply",)) or "").strip()
    if channel == "Direct Sales":
        return str(_row_value(data, ("client state",)) or "").strip()
    return ""


def _with_tax_amount(channel: str, row) -> float:
    aliases = {
        "DSG": ("order total amount",),
        "SFH": ("earnings",),
        "Direct Sales": ("total",),
    }
    return _number(_row_value(row.row_data, aliases[channel]))


def _order_identifier(channel: str, row) -> str:
    if channel == "SFH":
        value = _row_value(row.row_data, ("invoice no.", "receipt no.", "reg. no."))
    else:
        value = getattr(row, "order_number", None) or _row_value(
            row.row_data, ("order number", "invoice number")
        )
    text = str(value or "").strip().casefold()
    return text


def _style_sheet(sheet, title: str | None = None) -> None:
    sheet.freeze_panes = "A3" if title else "A2"
    header_row = 2 if title else 1
    if title:
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(sheet.max_column, 1))
        cell = sheet.cell(1, 1)
        cell.value = title
        cell.font = Font(color="FFFFFF", bold=True, size=14)
        cell.fill = TITLE_FILL
        cell.alignment = Alignment(horizontal="center")
    for cell in sheet[header_row]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    sheet.auto_filter.ref = f"A{header_row}:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"
    for column_index, column in enumerate(sheet.columns, 1):
        values = [str(cell.value or "") for cell in column]
        sheet.column_dimensions[get_column_letter(column_index)].width = min(max(map(len, values)) + 3, 45)


def _reconciled_whole_numbers(values: list[float], target: int | None = None) -> list[int]:
    """Round visible rows while preserving the rounded aggregate total."""
    displayed = [whole_number(value) for value in values]
    difference = (whole_number(sum(values)) if target is None else target) - sum(displayed)
    if difference == 0 or not values:
        return displayed
    errors = [value - rounded for value, rounded in zip(values, displayed)]
    direction = 1 if difference > 0 else -1
    order = sorted(
        range(len(values)),
        key=lambda index: errors[index],
        reverse=difference > 0,
    )
    for offset in range(abs(difference)):
        displayed[order[offset % len(order)]] += direction
    return displayed


def _report_year_month(month_key: str) -> tuple[int | str, str]:
    try:
        year, month = (int(value) for value in month_key.split("-"))
        return year, datetime(year, month, 1).strftime("%B")
    except (AttributeError, TypeError, ValueError):
        return "", ""


def _append_channel_sheet(
    workbook: Workbook, name: str, rows, label: str, *, include_period: bool = False
) -> None:
    sheet = workbook.create_sheet(_channel_display_name(name))
    sheet.append([label])
    sheet.append(
        ["Sl.No", "Year", "Month", "Type", "ID", "Description", "Quantity", "Total", "Location"]
        if include_period else
        ["Sl.No", "Type", "ID", "Description", "Quantity", "Total", "Location"]
    )
    display_quantities = _reconciled_whole_numbers(
        [_quantity(name, row) for _, row, _ in rows]
    )
    display_amounts = _reconciled_whole_numbers(
        [_amount(name, row) for _, row, _ in rows]
    )
    for index, ((_, row, month_key), quantity, amount) in enumerate(
        zip(rows, display_quantities, display_amounts), 1
    ):
        data = row.row_data
        values = [
            index,
            row.category or _row_value(data, ("type", "sales type")) or "",
            getattr(row, "order_number", None) or _row_value(data, ("id", "order id", "invoice number")) or "",
            row.product_name or getattr(row, "course", None) or "",
            quantity,
            amount,
            _location(name, data) or "NA",
        ]
        if include_period:
            year, month = _report_year_month(month_key)
            values[1:1] = [year, month]
        sheet.append(values)
    first_data_row = 3
    last_data_row = sheet.max_row
    has_data = last_data_row >= first_data_row
    if include_period:
        sheet.append([
            "", "", "", "Grand Total", "", "",
            f"=SUBTOTAL(109,G{first_data_row}:G{last_data_row})" if has_data else 0,
            f"=SUBTOTAL(109,H{first_data_row}:H{last_data_row})" if has_data else 0,
            "",
        ])
    else:
        sheet.append([
            "", "Grand Total", "", "",
            f"=SUBTOTAL(109,E{first_data_row}:E{last_data_row})" if has_data else 0,
            f"=SUBTOTAL(109,F{first_data_row}:F{last_data_row})" if has_data else 0,
            "",
        ])
    _style_sheet(sheet, label)
    sheet.auto_filter.ref = f"A2:{'I' if include_period else 'G'}{last_data_row if has_data else 2}"
    for column in (("G", "H") if include_period else ("E", "F")):
        for cell in sheet[column][2:]:
            _set_indian_whole_number_format(cell)
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8EAF6")
        cell.border = Border(bottom=Side(style="thin", color="1F2937"))


def _summary_taxable_value(basic_value: float, shipping: float, discount: float) -> float:
    return basic_value + shipping - discount


def _append_dsg_summary_sheet(workbook: Workbook, rows, label: str) -> None:
    """Build the Summary Report DSG detail sheet using order-level charges once."""
    sheet = workbook.create_sheet("DSG")
    sheet.append([label])
    sheet.append([
        "SL No", "Year", "Month", "Order ID", "Type", "Description", "Quantity",
        "Basic Value", "Shipping", "Discount", "Taxable Value",
        "Total Tax", "Total Invoice Value",
    ])
    completed_rows = [item for item in rows if _dsg_is_completed(item[1])]
    charged_orders: set[str] = set()
    for index, (_, row, month_key) in enumerate(completed_rows, 1):
        data = row.row_data
        year, month = _report_year_month(month_key)
        order_id = str(
            getattr(row, "order_number", None)
            or _row_value(data, ("order number",))
            or ""
        ).strip()
        order_key = order_id.casefold()
        first_order_row = order_key not in charged_orders
        if first_order_row:
            charged_orders.add(order_key)
        basic_value = _number(
            _row_value(data, ("item cost × quantity", "item cost x quantity"))
            if _row_value(data, ("item cost × quantity", "item cost x quantity")) is not None
            else row.amount
        )
        shipping = _number(_row_value(data, ("order shipping amount",))) if first_order_row else 0.0
        discount = _number(_row_value(data, ("cart discount amount",)))
        total_tax = _number(_row_value(data, ("order total tax amount",))) if first_order_row else 0.0
        taxable_value = _summary_taxable_value(basic_value, shipping, discount)
        sheet.append([
            index,
            year,
            month,
            order_id,
            row.category or _row_value(data, ("category",)) or "",
            row.product_name or _row_value(data, ("product name",)) or "",
            _quantity("DSG", row),
            basic_value,
            shipping,
            discount,
            taxable_value,
            total_tax,
            taxable_value + total_tax,
        ])
    first_data_row = 3
    last_data_row = sheet.max_row
    has_data = last_data_row >= first_data_row
    sheet.append([
        "", "", "", "Grand Total", "", "",
        f"=SUBTOTAL(109,G{first_data_row}:G{last_data_row})" if has_data else 0,
        *(
            [f"=SUBTOTAL(109,{column}{first_data_row}:{column}{last_data_row})" for column in "HIJKLM"]
            if has_data else [0] * 6
        ),
    ])
    _style_sheet(sheet, label)
    sheet.auto_filter.ref = f"A2:M{last_data_row if has_data else 2}"
    for column in "GHIJKLM":
        for cell in sheet[column][2:]:
            _set_indian_whole_number_format(cell)
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8EAF6")
        cell.border = Border(bottom=Side(style="thin", color="1F2937"))


def _append_sfh_summary_sheet(workbook: Workbook, rows, label: str) -> None:
    """Build one Summary Report detail row per unique SFH Invoice No."""
    sheet = workbook.create_sheet("SFH")
    sheet.append([label])
    sheet.append([
        "SL No", "Year", "Month", "Order ID", "Type", "Description", "Quantity",
        "Basic Value", "Shipping", "Discount", "Taxable Value",
        "Total Tax", "Total Invoice Value",
    ])
    seen_invoices: set[str] = set()
    output_rows = []
    for _, row, month_key in rows:
        data = row.row_data
        year, month = _report_year_month(month_key)
        invoice = str(
            _row_value(data, ("invoice no.", "invoice no", "invoice number")) or ""
        ).strip()
        invoice_key = invoice.casefold()
        if invoice_key in seen_invoices:
            continue
        seen_invoices.add(invoice_key)
        currency = _row_value(data, ("currency",))
        # SFH business rule: only the rupee symbol is INR. Text such as
        # "INR", and every other value, is treated as Non-INR.
        is_inr = sfh_is_inr_currency(currency)
        basic_value = _number(_row_value(
            data,
            ("without tax total",) if is_inr else ("earnings",),
        ))
        total_tax = _number(_row_value(data, ("tax",))) if is_inr else 0.0
        taxable_value = _summary_taxable_value(basic_value, 0, 0)
        output_rows.append([
            len(output_rows) + 1,
            year,
            month,
            invoice,
            "Web Version",
            row.course or row.product_name or _row_value(data, ("course",)) or "",
            1,
            basic_value,
            0,
            0,
            taxable_value,
            total_tax,
            taxable_value + total_tax,
        ])
    for output_row in output_rows:
        sheet.append(output_row)
    first_data_row = 3
    last_data_row = sheet.max_row
    has_data = last_data_row >= first_data_row
    sheet.append([
        "", "", "", "Grand Total", "", "",
        f"=SUBTOTAL(109,G{first_data_row}:G{last_data_row})" if has_data else 0,
        *(
            [f"=SUBTOTAL(109,{column}{first_data_row}:{column}{last_data_row})" for column in "HIJKLM"]
            if has_data else [0] * 6
        ),
    ])
    _style_sheet(sheet, label)
    sheet.auto_filter.ref = f"A2:M{last_data_row if has_data else 2}"
    for column in "GHIJKLM":
        for cell in sheet[column][2:]:
            _set_indian_whole_number_format(cell)
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8EAF6")
        cell.border = Border(bottom=Side(style="thin", color="1F2937"))


def _append_amazon_summary_sheet(workbook: Workbook, rows, label: str) -> None:
    """Build the Summary Report Amazon detail sheet without tax or discount."""
    sheet = workbook.create_sheet("Amazon")
    sheet.append([label])
    sheet.append([
        "SL No", "Year", "Month", "Order ID", "Type", "Description", "Quantity",
        "Basic Value", "Shipping", "Discount", "Taxable Value",
        "Total Tax", "Total Invoice Value",
    ])
    valid_rows = [item for item in rows if not _amazon_is_cancelled(item[1])]
    for index, (_, row, month_key) in enumerate(valid_rows, 1):
        data = row.row_data
        year, month = _report_year_month(month_key)
        basic_value = _number(_row_value(data, ("item-price", "item price")))
        shipping = _number(_row_value(data, ("shipping-price", "shipping price")))
        taxable_value = _summary_taxable_value(basic_value, shipping, 0)
        sheet.append([
            index,
            year,
            month,
            _row_value(data, ("amazon-order-id", "amazon order id")) or "",
            "Books",
            _row_value(data, ("product-name", "product name")) or row.product_name or "",
            _number(_row_value(data, ("quantity",))),
            basic_value,
            shipping,
            0,
            taxable_value,
            0,
            taxable_value,
        ])
    first_data_row = 3
    last_data_row = sheet.max_row
    has_data = last_data_row >= first_data_row
    sheet.append([
        "", "", "", "Grand Total", "", "",
        f"=SUBTOTAL(109,G{first_data_row}:G{last_data_row})" if has_data else 0,
        *(
            [f"=SUBTOTAL(109,{column}{first_data_row}:{column}{last_data_row})" for column in "HIJKLM"]
            if has_data else [0] * 6
        ),
    ])
    _style_sheet(sheet, label)
    sheet.auto_filter.ref = f"A2:M{last_data_row if has_data else 2}"
    for column in "GHIJKLM":
        for cell in sheet[column][2:]:
            _set_indian_whole_number_format(cell)
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8EAF6")
        cell.border = Border(bottom=Side(style="thin", color="1F2937"))


def _append_direct_sales_summary_sheet(workbook: Workbook, rows, label: str) -> None:
    """Build Direct Sales details from matched rows with invoice totals allocated once."""
    sheet = workbook.create_sheet("Direct Sales")
    sheet.append([label])
    sheet.append([
        "SL No", "Year", "Month", "Order ID", "Type", "Description", "Quantity",
        "Taxable Value", "Total Tax", "Total Invoice Value",
    ])

    invoice_totals: dict[str, float] = defaultdict(float)
    for _, row, _ in rows:
        invoice_key = _order_identifier("Direct Sales", row)
        invoice_totals[invoice_key] += _direct_amount(row)

    for index, (_, row, month_key) in enumerate(rows, 1):
        data = row.row_data
        year, month = _report_year_month(month_key)
        invoice_id = str(
            getattr(row, "order_number", None)
            or _row_value(data, ("invoice number",))
            or ""
        ).strip()
        invoice_key = invoice_id.casefold()
        # Upload processing allocates Without Tax Total over matched inventory
        # lines. Keep that allocation in the report so filtering a line does
        # not hide the entire invoice value. Allocate invoice tax by the same
        # ratio; the complete invoice still adds back exactly once.
        taxable_value = _direct_amount(row)
        invoice_taxable = invoice_totals.get(invoice_key, 0.0)
        invoice_tax = _number(_row_value(data, ("tax",)))
        total_tax = (
            invoice_tax * taxable_value / invoice_taxable
            if invoice_taxable else 0.0
        )
        sheet.append([
            index,
            year,
            month,
            invoice_id,
            row.category or _row_value(data, ("category",)) or "",
            row.product_name or _row_value(data, ("item details",)) or "",
            _number(_row_value(data, ("category quantity", "qty", "quantity"))),
            taxable_value,
            total_tax,
            taxable_value + total_tax,
        ])

    first_data_row = 3
    last_data_row = sheet.max_row
    has_data = last_data_row >= first_data_row
    sheet.append([
        "", "", "", "Grand Total", "", "",
        f"=SUBTOTAL(109,G{first_data_row}:G{last_data_row})" if has_data else 0,
        *(
            [f"=SUBTOTAL(109,{column}{first_data_row}:{column}{last_data_row})" for column in "HIJ"]
            if has_data else [0] * 3
        ),
    ])
    _style_sheet(sheet, label)
    sheet.auto_filter.ref = f"A2:J{last_data_row if has_data else 2}"
    for column in "GHIJ":
        for cell in sheet[column][2:]:
            _set_indian_whole_number_format(cell)
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8EAF6")
        cell.border = Border(bottom=Side(style="thin", color="1F2937"))


def _summary_gst(channel: str, rows) -> float:
    """Return GST using the same uniqueness/currency rules as detail sheets."""
    if channel == "Amazon":
        return 0.0
    seen_orders: set[str] = set()
    total = 0.0
    for _, row, _ in rows:
        data = row.row_data
        order_id = _order_identifier(channel, row)
        if order_id in seen_orders:
            continue
        seen_orders.add(order_id)
        if channel == "SFH":
            currency = _row_value(data, ("currency",))
            if not sfh_is_inr_currency(currency):
                continue
        total += _number(_row_value(data, ("order total tax amount", "tax")))
    return total


def _summary_workbook(rows, grain: str, period: str, year: int) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    label = _period_label(grain, period, year)
    metrics = _load_channel_metrics()
    summary = workbook.create_sheet("Summary")
    summary.append([f"{label} Sales and Out put GST -Summary"])
    summary.append([
        "Sl.no",
        "Particulars",
        "Orders",
        "Foreign Sales",
        "Books Sales",
        "Taxable Sales",
        "Total Sales",
        "GST Only On Taxable",
        "Total Invoice Amount",
        "Remarks",
    ])
    channel_definitions = (
        ("DSG", "dsg", "DSG"),
        ("SFH", "sfh", "SFH"),
        ("Direct Sales", "direct", "Direct Sales"),
        ("Amazon", "amazon", "Amazon Sales"),
    )
    raw_rows = []
    for channel_name, metric_key, display_name in channel_definitions:
        channel_rows = [item for item in rows if item[0] == channel_name]
        if not channel_rows:
            continue
        values = _for_period(metrics.get(metric_key, {}), grain, period, year) if metric_key in metrics else defaultdict(float)
        amounts = [
            _unique_order_count(channel_rows),
            values["zero_rated"],
            values["exempted"],
            values["taxable"],
            values["pnl"],
            _summary_gst(channel_name, channel_rows),
        ]
        raw_rows.append((len(raw_rows) + 1, display_name, amounts))
    # Reconcile each visible sales bucket, but round P&L from the combined raw
    # channel amount. This matches the detail-sheet grand total when fractional
    # values across multiple buckets combine to an additional rupee.
    reconciled_buckets = [
        _reconciled_whole_numbers([amounts[column] for _, _, amounts in raw_rows])
        for column in range(1, 4)
    ]
    displayed_rows = [
        [
            raw_rows[row][2][0],
            *(reconciled_buckets[column][row] for column in range(3)),
            whole_number(raw_rows[row][2][4]),
            whole_number(raw_rows[row][2][5]),
            whole_number(raw_rows[row][2][4]) + whole_number(raw_rows[row][2][5]),
            "",
        ]
        for row in range(len(raw_rows))
    ]
    for (index, display_name, _), display_amounts in zip(raw_rows, displayed_rows):
        summary.append([index, display_name, *display_amounts])
    bucket_totals = [sum(column) for column in reconciled_buckets]
    sales_total = sum(row[4] for row in displayed_rows)
    gst_total = sum(row[5] for row in displayed_rows)
    total_values = [sum(row[0] for row in displayed_rows), *bucket_totals, sales_total, gst_total, sales_total + gst_total, "-"]
    summary.append(["", "Total Sales", *total_values])
    _style_sheet(summary, f"{label} Sales and Out put GST -Summary")
    thin = Side(style="thin", color="1F2937")
    for row in summary.iter_rows(min_row=2, max_row=summary.max_row, min_col=1, max_col=10):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if 3 <= cell.column <= 9 and cell.row >= 3:
                _set_indian_whole_number_format(cell)
    for cell in summary[summary.max_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8EAF6")
    summary.row_dimensions[2].height = 34
    for channel, _ in CHANNELS:
        channel_rows = [item for item in rows if item[0] == channel]
        display_channel = channel
        channel_title = f"{label} {display_channel}" if channel == "Direct Sales" else f"{label} {display_channel} Sales"
        if channel == "DSG":
            _append_dsg_summary_sheet(workbook, channel_rows, channel_title)
        elif channel == "SFH":
            _append_sfh_summary_sheet(workbook, channel_rows, channel_title)
        elif channel == "Amazon":
            _append_amazon_summary_sheet(workbook, channel_rows, channel_title)
        elif channel == "Direct Sales":
            _append_direct_sales_summary_sheet(workbook, channel_rows, channel_title)
        else:
            _append_channel_sheet(workbook, channel, channel_rows, channel_title)
    return workbook


def _performance_workbook(rows, report_type: str, grain: str, period: str, year: int) -> Workbook:
    if report_type == "channel":
        return _channel_performance_workbook(rows)
    if report_type == "category":
        return _category_performance_workbook(rows, grain, period, year)
    if report_type == "product":
        return _product_performance_workbook(rows)
    workbook = Workbook()
    sheet = workbook.active
    titles = {"channel": "Channel Wise Performance", "category": "Category Wise Performance", "product": "Product Wise Performance"}
    label = _period_label(grain, period, year)
    sheet.title = titles[report_type][:31]
    sheet.append([f"{label} {titles[report_type]}"])
    group_index = {"channel": 0, "category": 1, "product": 2}[report_type]
    totals = defaultdict(lambda: [0.0, 0.0])
    for channel, row, _ in rows:
        keys = (_channel_display_name(channel), row.category or "Uncategorised", row.product_name or getattr(row, "course", None) or "Unmapped")
        quantity = _quantity(channel, row)
        totals[keys[group_index]][0] += quantity
        totals[keys[group_index]][1] += _amount(channel, row)
    sheet.append(["Sl.No", titles[report_type].replace(" Wise Performance", ""), "Quantity", "Total"])
    for index, (name, values) in enumerate(sorted(totals.items()), 1):
        sheet.append([index, name, whole_number(values[0]), whole_number(values[1])])
    _style_sheet(sheet, f"{label} {titles[report_type]}")
    for column in ("C", "D"):
        for cell in sheet[column][2:]:
            _set_indian_whole_number_format(cell)
    return workbook


def _product_type(channel: str, row) -> str:
    """Return the source category used as Type in product performance."""
    if channel == "SFH":
        return "Web Version"
    return str(
        getattr(row, "category", None)
        or _row_value(row.row_data, ("category",))
        or "Uncategorised"
    ).strip()


def _product_performance_workbook(rows) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    grouped = defaultdict(
        lambda: {"order_ids": set(), "types": set(), "quantity": 0.0, "total": 0.0}
    )
    for channel, row, month_key in rows:
        product = str(row.product_name or getattr(row, "course", None) or "Unmapped").strip()
        key = (month_key, channel, product)
        grouped[key]["types"].add(_product_type(channel, row))
        order_id = _order_identifier(channel, row)
        if order_id:
            grouped[key]["order_ids"].add(order_id)
        grouped[key]["quantity"] += _quantity(channel, row)
        grouped[key]["total"] += _amount(channel, row)
    for channel, _ in CHANNELS:
        sheet = workbook.create_sheet(channel)
        sheet.append([
            "Year", "Month", "Channel", "Type", "Product", "Orders",
            "Quantity", "Without Tax Total",
        ])
        ordered = sorted(
            (
                (key, values)
                for key, values in grouped.items()
                if key[1] == channel
            ),
            key=lambda item: (item[0][0], item[0][2].casefold()),
        )
        display_quantities = [0] * len(ordered)
        display_totals = [0] * len(ordered)
        month_groups = defaultdict(list)
        for index, ((month_key, _, _), _) in enumerate(ordered):
            month_groups[month_key].append(index)
        for indexes in month_groups.values():
            reconciled_quantities = _reconciled_whole_numbers(
                [ordered[index][1]["quantity"] for index in indexes]
            )
            reconciled_totals = _reconciled_whole_numbers(
                [ordered[index][1]["total"] for index in indexes]
            )
            for position, index in enumerate(indexes):
                display_quantities[index] = reconciled_quantities[position]
                display_totals[index] = reconciled_totals[position]
        for ((month_key, _, product), values), quantity, total in zip(
            ordered, display_quantities, display_totals
        ):
            row_year, row_month = (int(value) for value in month_key.split("-"))
            sheet.append([
                row_year,
                datetime(row_year, row_month, 1).strftime("%b"),
                channel,
                ", ".join(sorted(values["types"], key=str.casefold)),
                product,
                len(values["order_ids"]),
                quantity,
                total,
            ])
        first_data_row = 2
        last_data_row = sheet.max_row
        has_data = last_data_row >= first_data_row
        total_orders = (
            f"=SUBTOTAL(109,F{first_data_row}:F{last_data_row})" if has_data else 0
        )
        total_quantity = (
            f"=SUBTOTAL(109,G{first_data_row}:G{last_data_row})" if has_data else 0
        )
        total_without_tax = (
            f"=SUBTOTAL(109,H{first_data_row}:H{last_data_row})" if has_data else 0
        )
        sheet.append([
            "", "", channel, "", "Grand Total", total_orders,
            total_quantity, total_without_tax,
        ])
        _style_sheet(sheet)
        # Keep Grand Total outside the AutoFilter so it remains visible. SUBTOTAL
        # recalculates against only the rows visible after an Excel filter.
        sheet.auto_filter.ref = f"A1:H{last_data_row if has_data else 1}"
        for row_number in range(2, sheet.max_row + 1):
            for column in (6, 7, 8):
                _set_indian_whole_number_format(sheet.cell(row_number, column))
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E8EAF6")
    workbook.properties.title = "Product Wise Performance Report"
    return workbook


def _category_performance_workbook(rows, grain: str, period: str, year: int) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    if grain == "monthly":
        month_numbers = sorted({int(value) for value in period.split(",")})
    else:
        month_numbers = sorted({int(month_key.split("-")[1]) for _, _, month_key in rows})
    month_keys = [f"{year}-{month:02d}" for month in month_numbers]
    month_labels = [datetime(year, month, 1).strftime("%b") for month in month_numbers]
    canonical = {
        "audio device": "Audio Device",
        "books": "Books",
        "book": "Books",
        "pen drive": "Pen Drive",
        "pen drives": "Pen Drive",
        "web version": "Web Version",
        "web / e-books": "Web Version",
        "web version / e-books": "Web Version",
    }
    preferred_order = ["Audio Device", "Books", "Pen Drive", "Web Version"]
    for channel, _ in CHANNELS:
        sheet = workbook.create_sheet(channel)
        sheet.append(["Category", *month_labels, "Total Sales", "% Contribution"])
        monthly_totals = defaultdict(lambda: defaultdict(float))
        for row_channel, row, month_key in rows:
            if row_channel != channel:
                continue
            raw_category = str(row.category or "Uncategorised").strip()
            category = canonical.get(raw_category.casefold(), raw_category or "Uncategorised")
            monthly_totals[category][month_key] += _amount(channel, row)
        categories = [category for category in preferred_order if category in monthly_totals]
        categories.extend(sorted(
            (category for category in monthly_totals if category not in preferred_order),
            key=str.casefold,
        ))
        displayed_by_category = {category: [0] * len(month_keys) for category in categories}
        for month_index, month_key in enumerate(month_keys):
            reconciled = _reconciled_whole_numbers(
                [monthly_totals[category][month_key] for category in categories]
            )
            for category_index, category in enumerate(categories):
                displayed_by_category[category][month_index] = reconciled[category_index]
        category_values = [
            (category, values, sum(values))
            for category, values in displayed_by_category.items()
        ]
        grand_total = sum(total for _, _, total in category_values)
        for category, values, total in category_values:
            contribution = total / grand_total if grand_total else 0
            sheet.append([category, *values, total, contribution])
        monthly_grand_totals = [sum(values[index] for _, values, _ in category_values) for index in range(len(month_keys))]
        sheet.append(["Total", *monthly_grand_totals, grand_total, "-"])
        _style_sheet(sheet)
        total_column = len(month_keys) + 2
        contribution_column = total_column + 1
        for row_number in range(2, sheet.max_row + 1):
            for column in range(2, total_column + 1):
                _set_indian_whole_number_format(sheet.cell(row_number, column))
            if isinstance(sheet.cell(row_number, contribution_column).value, (int, float)):
                sheet.cell(row_number, contribution_column).number_format = '0%'
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E8EAF6")
    workbook.properties.title = "Category Wise Performance Report"
    return workbook


def _channel_performance_workbook(rows) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    headers = [
        "Year", "Month Name", "Total Orders", "Orders Var MOM", "Orders MOM %",
        "With Tax Total", "Without Tax Total", "Sales Var MOM", "Sales MOM %",
    ]
    for channel, _ in CHANNELS:
        sheet = workbook.create_sheet(channel)
        sheet.append(headers)
        monthly = defaultdict(lambda: {"with_tax_by_order": defaultdict(set), "without_tax": 0.0})
        for row_channel, row, month_key in rows:
            if row_channel != channel:
                continue
            if channel == "Amazon":
                order_status = _row_value(row.row_data, ("order-status", "order status"))
                if str(order_status or "").strip().casefold() != "shipped - delivered to buyer":
                    continue
                order_id = str(
                    _row_value(row.row_data, ("amazon-order-id", "amazon order id")) or ""
                ).strip().casefold()
            else:
                order_id = _order_identifier(channel, row)
            if order_id:
                monthly[month_key]["with_tax_by_order"][order_id]
            if channel == "Amazon":
                monthly[month_key].setdefault("amazon_with_tax", 0.0)
                monthly[month_key]["amazon_with_tax"] += _number(
                    _row_value(row.row_data, ("item-price", "item price"))
                )
            elif order_id:
                monthly[month_key]["with_tax_by_order"][order_id].add(
                    _with_tax_amount(channel, row)
                )
            monthly[month_key]["without_tax"] += _amount(channel, row)
        previous = None
        for month_key in sorted(monthly):
            values = monthly[month_key]
            row_year, row_month = (int(value) for value in month_key.split("-"))
            orders = len(values["with_tax_by_order"])
            with_tax = whole_number(
                values.get("amazon_with_tax", 0.0)
                if channel == "Amazon"
                else sum(sum(order_values) for order_values in values["with_tax_by_order"].values())
            )
            without_tax = whole_number(values["without_tax"])
            if previous is None:
                order_variance = order_percent = sales_variance = sales_percent = "-"
            else:
                order_variance = orders - previous["orders"]
                order_percent = (
                    order_variance / previous["orders"] if previous["orders"]
                    else "-"
                ) if channel == "Amazon" else order_variance / orders if orders else "-"
                sales_variance = without_tax - previous["without_tax"]
                sales_percent = sales_variance / without_tax if without_tax else "-"
            sheet.append([
                row_year,
                datetime(row_year, row_month, 1).strftime("%b"),
                orders,
                order_variance,
                order_percent,
                with_tax,
                without_tax,
                sales_variance,
                sales_percent,
            ])
            previous = {
                "orders": orders,
                "with_tax": with_tax,
                "without_tax": without_tax,
            }
        _style_sheet(sheet)
        for row_number in range(2, sheet.max_row + 1):
            for column in (3, 4, 6, 7, 8):
                _set_indian_whole_number_format(sheet.cell(row_number, column))
            for column in (5, 9):
                if isinstance(sheet.cell(row_number, column).value, (int, float)):
                    sheet.cell(row_number, column).number_format = '0%;[Red]-0%;-'
    workbook.properties.title = "Channel Wise Performance Report"
    return workbook


def _raw_workbook(rows, grain: str, period: str, year: int) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    label = _period_label(grain, period, year)
    for channel, _ in CHANNELS:
        channel_rows = [row for row_channel, row, _ in rows if row_channel == channel]
        sheet = workbook.create_sheet(channel)
        keys = list(dict.fromkeys(key for row in channel_rows for key in row.row_data))
        sheet.append(keys or ["No records"])
        for row in channel_rows:
            sheet.append([row.row_data.get(key, "") for key in keys])
        _style_sheet(sheet)
    workbook.properties.title = f"{label} filtered data set"
    return workbook


def _clean_workbook(rows, grain: str, period: str, year: int) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    label = _period_label(grain, period, year)
    for channel, _ in CHANNELS:
        channel_rows = [item for item in rows if item[0] == channel]
        display_channel = _channel_display_name(channel)
        channel_title = f"{label} {display_channel}" if channel == "Direct Sales" else f"{label} {display_channel} Sales"
        _append_channel_sheet(workbook, channel, channel_rows, channel_title)
    workbook.properties.title = f"{label} cleaned sales dataset"
    return workbook


@router.get("/direct-sales-overview/options")
def direct_sales_overview_options() -> dict[str, object]:
    with SessionLocal() as database:
        source = _direct_overview_source(database)
    return {
        "years": sorted({item["year"] for item in source}, reverse=True),
        "months": sorted({item["month"] for item in source}),
        "types": list(DIRECT_OVERVIEW_TYPE_ORDER),
        "products": sorted({item["product"] for item in source}, key=str.casefold),
    }


@router.get("/direct-sales-overview/preview")
def preview_direct_sales_overview(
    years: list[int] = Query(...),
    months: list[int] = Query(...),
    types: list[str] = Query(...),
    products: list[str] | None = Query(None),
) -> dict[str, object]:
    with SessionLocal() as database:
        result = _build_direct_sales_overview(
            database,
            years=set(years),
            months=set(months),
            types=set(types),
            products=set(products or []),
        )
    return {key: value for key, value in result.items() if key != "rows"}


@router.get("/direct-sales-overview/download")
def download_direct_sales_overview(
    years: list[int] = Query(...),
    months: list[int] = Query(...),
    types: list[str] = Query(...),
    products: list[str] | None = Query(None),
):
    with SessionLocal() as database:
        result = _build_direct_sales_overview(
            database,
            years=set(years),
            months=set(months),
            types=set(types),
            products=set(products or []),
        )
    output = StringIO()
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(["Year", "Month", "Type", "Product", "Orders", "Invoice Qty (Classification)", "Product Quantity", "Without Tax Total"])
    for row in result["rows"]:
        writer.writerow([
            row["year"], row["month"], row["type"], row["product"],
            _indian_integer(row["orders"]),
            _indian_integer(row["invoice_quantity"]),
            _indian_integer(row["quantity"]),
            _indian_integer(row["without_tax_total"]),
        ])
    filename = f"direct-sales-overview-{'-'.join(map(str, sorted(set(years))))}.csv"
    return Response(
        content="\ufeff" + output.getvalue(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/download")
def download_report(
    grain: str = "monthly",
    period: str | None = None,
    year: int | None = None,
    report_type: str = "summary",
    dataset: str = "report",
):
    now = datetime.now()
    year = year or now.year
    period = period or (str(now.month) if grain == "monthly" else str(year))
    if grain not in {"monthly", "yearly"} or report_type not in REPORT_TYPES or dataset not in {"report", "clean"}:
        raise HTTPException(status_code=422, detail="Invalid report filter.")
    try:
        if grain == "monthly":
            selected_months = [int(value) for value in period.split(",") if value]
            if not selected_months or any(value < 1 or value > 12 for value in selected_months):
                raise ValueError
    except ValueError as error:
        raise HTTPException(status_code=422, detail="Invalid reporting period.") from error
    with SessionLocal() as database:
        rows = _filtered_rows(database, grain, period, year)
    workbook = _clean_workbook(rows, grain, period, year) if dataset == "clean" else (
        _summary_workbook(rows, grain, period, year)
        if report_type == "summary"
        else _performance_workbook(rows, report_type, grain, period, year)
    )
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    label = _period_label(grain, period, year).replace(" ", "-")
    filename = f"{label}-{report_type}-{dataset}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )
