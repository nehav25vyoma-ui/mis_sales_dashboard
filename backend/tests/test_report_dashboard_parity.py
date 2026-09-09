from datetime import datetime
from math import floor
from types import SimpleNamespace
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app import dashboard, reports
from app.database.models import DirectSalesDatasetRow, DSGDatasetRow, SFHDatasetRow, AmazonDatasetRow
from test_reports import _direct_row, _OverviewDatabase, _dsg_summary_row, _sfh_summary_row, _amazon_summary_row


@pytest.fixture
def source(monkeypatch):
    rows = [
        _direct_row(1, "A", "Book", 100.4, 1),
        _direct_row(2, "B", "book", 50.4, 1),
        _direct_row(3, "C", "Language Lab", 25.4, 1, "Language Lab"),
        _direct_row(4, "D", "Unknown", 10.4, 1),
        _direct_row(5, "E", "Stall", 20.4, 11, "stall"),
    ]
    rows[2].category = "Language Lab"
    rows[3].category = "N/A"
    monkeypatch.setattr(dashboard, "_dashboard_source_cache", {"version": "parity"})
    monkeypatch.setattr(dashboard, "_channel_metrics_cache", {})
    monkeypatch.setattr(dashboard, "_cached_rows", lambda model: rows if model is DirectSalesDatasetRow else [])
    monkeypatch.setattr(dashboard, "_cached_upload_dates", lambda: {"upload-1": datetime(2026, 4, 1)})
    monkeypatch.setattr(dashboard, "saved_plan_years", lambda: set())
    monkeypatch.setattr(dashboard, "category_plans_for_year", lambda _: dict.fromkeys(dashboard.CATEGORY_MONTHLY_PLANS, 0))
    monkeypatch.setattr(dashboard, "plans_for_year", lambda _: {})
    monkeypatch.setattr(reports, "dashboard_kpis", dashboard._build_dashboard_kpis)
    return rows


def visual(channel="direct", period="4", view="overview"):
    return dashboard._build_dashboard_kpis(channel=channel, grain="monthly", period=period,
                                          year=2026, comparison_grain="monthly",
                                          comparison_period="3", comparison_year=2026, view=view)


def test_summary_and_channel_sales_match_their_visuals(source):
    rows = [("Direct Sales", row, "2026-04") for row in source]
    data = visual()
    cards = {card["id"]: card for card in data["cards"]}
    sheet = reports._summary_workbook(rows, "monthly", "4", 2026)["Summary"]
    for column, metric in ((4, "zero_rated"), (5, "exempted"), (6, "taxable"), (7, "pnl")):
        assert sheet.cell(3, column).value == floor(cards[metric]["total"] + 0.5)
    channel = reports._channel_performance_workbook(rows)["Direct Sales"]
    assert channel.cell(2, 7).value == floor(data["direct_sales_performance"]["current"]["Total Direct Sales"] + 0.5)


def test_category_report_uses_visual_values_and_selected_period_total(source):
    source[1].row_data["Issue Date"] = "2026-05-10"
    rows = [("Direct Sales", row, dashboard._month(row.row_data, datetime(2026, 4, 1))) for row in source]
    sheet = reports._category_performance_workbook(rows, "monthly", "4,5", 2026)["Direct Sales"]
    labels = {"Books": "Books", "Audio Device": "Audio Device", "Pen Drive": "Pen Drives", "Web Version": "Web / E-Books"}
    expected = reports._category_visual_values("direct", "monthly", "4,5", 2026)
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row - 1):
        assert row[3].value == expected[labels[row[0].value]]
    assert sheet.cell(sheet.max_row, 4).value == floor(visual(period="4,5")["channel_wise_performance"]["current"]["Total Sales"] + 0.5)


def test_product_report_preserves_precision_and_case_insensitive_grouping(source):
    rows = [("Direct Sales", row, "2026-04") for row in source]
    sheet = reports._product_performance_workbook(rows)["Direct Sales"]
    products = {sheet.cell(index, 5).value.casefold(): sheet.cell(index, 8).value
                for index in range(2, sheet.max_row)}
    data = visual(view="product")
    totals = {}
    for row in data["product_performance"]["details"]:
        name = row["description"].casefold()
        totals[name] = totals.get(name, 0) + row["total_invoice_value"]
    assert products == pytest.approx(totals)
    assert products["book"] == pytest.approx(150.8)


def test_overview_classifications_match_dashboard_without_rounding_transfers(source):
    result = reports._build_direct_sales_overview(
        _OverviewDatabase(source, [SimpleNamespace(upload_id="upload-1", uploaded_at=datetime(2026, 4, 1))]),
        years={2026}, months={4}, types=set(reports.DIRECT_OVERVIEW_TYPE_ORDER), products=set(),
    )
    totals = visual()["direct_sales_performance"]["current"]
    for name in reports.DIRECT_OVERVIEW_TYPE_ORDER:
        assert result["totals"][name] == floor(totals[name] + 0.5)
    assert sum(row["without_tax_total"] for row in result["rows"]) == pytest.approx(sum(dashboard._direct_amount(row) for row in source))


def test_vedanta_retail_is_counted_in_dashboard_and_overview(source):
    source[0].row_data["Private Notes"] = "Retail sales Vedanta book house"
    totals = visual()["direct_sales_performance"]["current"]
    assert totals["Retail"] == 100.4
    result = reports._build_direct_sales_overview(
        _OverviewDatabase(source, [SimpleNamespace(upload_id="upload-1", uploaded_at=datetime(2026, 4, 1))]),
        years={2026}, months={4}, types={"Retail"}, products=set(),
    )
    assert result["source_records"] == 1
    assert result["totals"]["Retail"] == floor(totals["Retail"] + 0.5)


def test_vedantha_client_with_blank_notes_reaches_retail_in_existing_records(source):
    source[0].row_data["Client Name"] = "Vedantha Book House"
    source[0].row_data["Private Notes"] = ""
    source[0].row_data["Category Quantity"] = 25
    assert visual()["direct_sales_performance"]["current"]["Retail"] == 100.4
    assert reports._direct_overview_type(source[0]) == "Retail"


def test_bulk_private_notes_rule_matches_visual_and_report_for_saved_rows(source):
    source[0].row_data.update({"Private Notes": "BULK sale", "Category Quantity": 1,
                               "Sales Classification": "In Office"})
    source[1].row_data.update({"Private Notes": "", "Category Quantity": 99,
                               "Sales Classification": "Bulk"})
    totals = visual()["channel_wise_performance"]["current"]
    assert totals["Bulk"] == 100.4
    assert reports._direct_overview_type(source[0]) == "Bulk"
    assert reports._direct_overview_type(source[1]) == "In Office"
    result = reports._build_direct_sales_overview(
        _OverviewDatabase(source, [SimpleNamespace(upload_id="upload-1", uploaded_at=datetime(2026, 4, 1))]),
        years={2026}, months={4}, types={"Bulk"}, products=set(),
    )
    assert result["source_records"] == 1
    assert result["totals"]["Bulk"] == floor(totals["Bulk"] + 0.5)


@pytest.mark.parametrize("notes,product,client,expected", [
    ("  STALL event  ", "Book", "", "Stall"),
    ("", "Stall Book", "", "In Office"),
    ("", "Vedanta Book House", "", "In Office"),
    ("", "Book", "  VEDANTA BOOK HOUSE  ", "Retail"),
    ("Retail sales Vedantha Book House", "Book", "", "Retail"),
    (" PHONE order ", "Book", "", "Call"),
    ("In office by Phaneendra", "Book", "", "In Office"),
    ("in office bulk", "Book", "", "In Office"),
    (" COURSE   PROMOTION bulk ", "Book", "", "Course Promotion"),
    ("", "S101 Grammer Course", "", "In Office"),
    ("Bulk order", "Book", "", "Bulk"),
    ("in office course promotion", "Language Lab", "", "Language Lab"),
])
def test_invoice_field_subtype_rules(notes, product, client, expected):
    row = _direct_row(1, "1", product, 100, 99, notes)
    row.row_data["Client Name"] = client
    assert dashboard._direct_sales_channel(row) == expected
    assert reports._direct_overview_type(row) == expected


def test_all_six_channel_performance_totals_use_invoice_rules(source):
    source.clear()
    cases = [("stall", "", "Stall"), ("", "Vedanta Book House", "Retail"),
             ("phone", "", "Call"), ("in office", "", "In Office"),
             ("course promotion", "", "Course Promotion"), ("bulk", "", "Bulk")]
    expected = {}
    for index, (notes, client, subtype) in enumerate(cases, 1):
        row = _direct_row(index, str(index), "Book", index * 100, 1, notes)
        row.row_data["Client Name"] = client
        source.append(row)
        expected[subtype] = index * 100
    for channel in ("all", "direct"):
        values = visual(channel)["channel_wise_performance"]["current"]
        assert {name: values[name] for name in expected} == expected
        assert values["Total Sales"] == sum(expected.values())


@pytest.mark.parametrize("channel", ["all", "direct"])
def test_channel_wise_performance_retail_in_both_periods(source, channel):
    for row in source[:2]:
        row.row_data["Client Name"] = "Vedantha Book House"
        row.row_data["Private Notes"] = ""
    source[1].row_data["Issue Date"] = "2026-05-10"
    data = dashboard._build_dashboard_kpis(
        channel=channel, grain="monthly", period="4", year=2026,
        comparison_grain="monthly", comparison_period="5", comparison_year=2026,
    )
    values = data["channel_wise_performance"]
    assert values["current"]["Retail"] == 100.4
    assert values["comparison"]["Retail"] == 50.4
    for period in ("current", "comparison"):
        assert values[period]["Total Sales"] == pytest.approx(sum(
            values[period][name] for name in (
                "Digital Online", "In Office", "Stall", "Bulk", "Call", "Retail", "Course Promotion",
            )
        ))


def test_yearly_product_filter_uses_the_same_fiscal_period(monkeypatch):
    class January(datetime):
        @classmethod
        def now(cls):
            return cls(2026, 1, 15)
    monkeypatch.setattr(dashboard, "datetime", January)
    for month, included in (("2025-03", False), ("2025-04", True), ("2025-12", True), ("2026-01", True), ("2026-02", False)):
        assert reports._performance_month_matches(month, "yearly", "2026", 2026) is included


def test_all_channel_summary_rounding_matches_kpi_and_workbooks_serialize(source, monkeypatch):
    dsg = _dsg_summary_row("D", "Completed", "Book", 90.49, 1, 0, 0, 0)
    sfh = _sfh_summary_row("S", "USD", "Course", 999, 80.49, 0)
    amazon = _amazon_summary_row("A", "Shipped - Delivered to Buyer", "Book", 1, 70.49, 0)
    grouped = {DSGDatasetRow: [dsg], SFHDatasetRow: [sfh], AmazonDatasetRow: [amazon], DirectSalesDatasetRow: source}
    for index, row in enumerate((dsg, sfh, amazon)):
        row.upload_id = "upload-1"
        row.order_number = str(index)
        row.currency = "INR"
        row.id = index
    monkeypatch.setattr(dashboard, "_cached_rows", lambda model: grouped.get(model, []))
    rows = [(channel, row, "2026-04") for channel, model in reports.CHANNELS for row in grouped[model]]
    workbook = reports._summary_workbook(rows, "monthly", "4", 2026)
    card = next(card for card in visual("all")["cards"] if card["id"] == "pnl")
    expected = [floor(item["value"] + 0.5) for item in card["breakdown"]]
    largest = max(range(len(expected)), key=lambda index: abs(expected[index]))
    expected[largest] += floor(card["total"] + 0.5) - sum(expected)
    by_channel = {item["label"]: value for item, value in zip(card["breakdown"], expected)}
    summary = workbook["Summary"]
    for row_index in range(3, summary.max_row):
        label = summary.cell(row_index, 2).value.replace("Amazon Sales", "Amazon")
        assert summary.cell(row_index, 7).value == by_channel[label]
    assert summary.cell(summary.max_row, 7).value == floor(card["total"] + 0.5)
    for output in (workbook, reports._product_performance_workbook(rows), reports._channel_performance_workbook(rows),
                   reports._category_performance_workbook(rows, "monthly", "4", 2026)):
        stream = BytesIO()
        output.save(stream)
        stream.seek(0)
        reopened = load_workbook(stream)
        assert reopened.sheetnames == output.sheetnames
