from types import SimpleNamespace
from datetime import datetime

import pytest

from app.reports import (
    INDIAN_LAKH_FORMAT,
    INDIAN_WHOLE_NUMBER_FORMAT,
    _append_channel_sheet,
    _append_dsg_summary_sheet,
    _append_sfh_summary_sheet,
    _append_amazon_summary_sheet,
    _append_direct_sales_summary_sheet,
    _build_direct_sales_overview,
    _direct_overview_type,
    _channel_performance_workbook,
    _product_performance_workbook,
    _product_type,
    _summary_workbook,
)
from app.database.models import AmazonDatasetRow, DirectSalesDatasetRow, DSGDatasetRow, SFHDatasetRow, UploadHistory
from app.dashboard import _amazon_is_cancelled, _amazon_is_excluded, _direct_amount, _direct_sales_channel, _dsg_basic_amount, _dsg_pnl_amount, _financial_breakdown, _month, _unique_order_count
from openpyxl import Workbook


def _row(category: str | None):
    return SimpleNamespace(category=category, row_data={"Category": category})


def test_product_type_uses_channel_category_rules():
    assert _product_type("DSG", _row("Books")) == "Books"
    assert _product_type("SFH", _row("Web Version")) == "Web Version"
    assert _product_type("Direct Sales", _row("Audio Device")) == "Audio Device"


def test_product_report_has_type_column_and_indian_number_format():
    workbook = _product_performance_workbook([])
    sheet = workbook["DSG"]

    assert [cell.value for cell in sheet[1]] == [
        "Year", "Month", "Channel", "Type", "Product", "Orders",
        "Quantity", "Without Tax Total",
    ]
    assert sheet.cell(2, 8).number_format != "General"


def test_product_report_grand_total_updates_with_excel_filters():
    row = SimpleNamespace(
        product_name="Spoken English",
        course=None,
        category="Books",
        order_number="DSG-1",
        amount="1234567",
        row_data={"Quantity": 2},
    )
    workbook = _product_performance_workbook([("DSG", row, "2026-08")])
    sheet = workbook["DSG"]

    assert sheet.auto_filter.ref == "A1:H2"
    assert sheet.cell(3, 5).value == "Grand Total"
    assert sheet.cell(3, 6).value == "=SUBTOTAL(109,F2:F2)"
    assert sheet.cell(3, 7).value == "=SUBTOTAL(109,G2:G2)"
    assert sheet.cell(3, 8).value == "=SUBTOTAL(109,H2:H2)"
    assert sheet.cell(2, 8).number_format == INDIAN_LAKH_FORMAT
    assert sheet.cell(3, 8).number_format == INDIAN_WHOLE_NUMBER_FORMAT


def test_product_report_counts_unique_amazon_order_ids():
    def amazon_row(order_id: str):
        return SimpleNamespace(
            product_name="Book A",
            course=None,
            category="Books",
            amount="100",
            row_data={"amazon-order-id": order_id, "quantity": 1},
        )

    workbook = _product_performance_workbook([
        ("Amazon", amazon_row("AMZ-1"), "2026-08"),
        ("Amazon", amazon_row("AMZ-1"), "2026-08"),
        ("Amazon", amazon_row("AMZ-2"), "2026-08"),
    ])
    sheet = workbook["Amazon"]

    assert sheet.cell(2, 6).value == 2
    assert sheet.cell(3, 6).value == "=SUBTOTAL(109,F2:F2)"


def test_summary_displayed_pnl_equals_displayed_sales_buckets(monkeypatch):
    metrics = {
        "dsg": {"zero_rated": 14234.77, "exempted": 36558.0, "taxable": 7515.0, "pnl": 58307.77},
        "sfh": {"zero_rated": 2314.35, "exempted": 0.0, "taxable": 8592.372881, "pnl": 10906.722881},
        "direct": {"zero_rated": 0.0, "exempted": 94244.833336, "taxable": 17816.736664, "pnl": 112061.57},
    }
    monkeypatch.setattr("app.reports._load_channel_metrics", lambda: metrics)
    monkeypatch.setattr("app.reports._for_period", lambda values, *_: values)

    summary_rows = [
        ("DSG", _dsg_summary_row("D-1", "Completed", "Book", 1, 1, 0, 0, 0), "2026-07"),
        ("SFH", _sfh_summary_row("S-1", "USD", "Course", 1, 1, 0), "2026-07"),
        ("Direct Sales", _direct_row(1, "X-1", "Book", 1, 1), "2026-07"),
    ]
    sheet = _summary_workbook(summary_rows, "monthly", "7", 2026)["Summary"]

    for row_number in range(3, 7):
        buckets = sum(sheet.cell(row_number, column).value for column in (4, 5, 6))
        assert abs(sheet.cell(row_number, 7).value - buckets) <= 1
        assert sheet.cell(row_number, 9).value == sheet.cell(row_number, 7).value + sheet.cell(row_number, 8).value
    assert sheet.cell(5, 7).value == 1
    assert [sheet.cell(row, 3).value for row in range(3, 6)] == [1, 1, 1]
    assert [sheet.cell(row, 2).value for row in range(3, 6)] == [
        "DSG", "SFH", "Direct Sales",
    ]


def test_summary_direct_sales_matches_detail_rows_and_includes_na(monkeypatch):
    monkeypatch.setattr("app.reports._load_channel_metrics", lambda: {
        "direct": {
            "zero_rated": 0.0, "exempted": 9999.0,
            "taxable": 9999.0, "pnl": 19998.0,
        },
    })
    monkeypatch.setattr("app.reports._for_period", lambda values, *_: values)
    rows = [
        ("Direct Sales", _direct_row(1, "D-1", "Book", 100.4, 1), "2026-07"),
        ("Direct Sales", SimpleNamespace(
            id=2, upload_id="upload", source_row_number=2,
            order_number="D-2", category="N/A", product_name="Other",
            amount="50.4", row_data={
                "Invoice Number": "D-2", "Without Tax Total": 50.4,
                "Category": "N/A", "Item Details": "Other", "Quantity": 1,
            },
        ), "2026-07"),
    ]

    workbook = _summary_workbook(rows, "monthly", "7", 2026)
    summary = workbook["Summary"]
    detail = workbook["Direct Sales"]

    assert summary.cell(3, 5).value == 100
    assert summary.cell(3, 6).value == 50
    assert summary.cell(3, 7).value == 151
    assert sum(detail.cell(row, 8).value for row in (3, 4)) == pytest.approx(150.8)


def test_clean_report_total_matches_rounded_raw_aggregate():
    rows = []
    for index, amount in enumerate((0.51, 0.51), 1):
        row = SimpleNamespace(
            category="Books",
            order_number=f"DSG-{index}",
            product_name="Book",
            amount=str(amount),
            row_data={"Quantity": 1, "Country Code (Billing)": "IN"},
        )
        rows.append(("DSG", row, "2026-04"))

    workbook = Workbook()
    workbook.remove(workbook.active)
    _append_channel_sheet(workbook, "DSG", rows, "Apr - 26 DSG Sales")
    sheet = workbook["DSG"]

    # Reconcile visible whole rows so their manual sum equals the rounded raw
    # aggregate instead of independently rounding both values up to 2.
    assert sheet.cell(3, 6).value == 0
    assert sheet.cell(4, 6).value == 1
    assert sheet.cell(sheet.max_row, 6).value == "=SUBTOTAL(109,F3:F4)"
    assert sheet.auto_filter.ref == "A2:G4"


def test_month_uses_the_current_row_date_without_identity_cache():
    row_data = {"Order Date": "2026-04-15"}

    assert _month(row_data, datetime(2026, 1, 1)) == "2026-04"
    row_data["Order Date"] = "2026-05-15"
    assert _month(row_data, datetime(2026, 1, 1)) == "2026-05"


def _dsg_summary_row(order, status, product, basic, quantity, shipping, discount, tax, order_total=None):
    return SimpleNamespace(
        order_number=order,
        category="Books",
        product_name=product,
        amount=str(basic),
        row_data={
            "Order Number": order,
            "Order Status": status,
            "Category": "Books",
            "Product Name": product,
            "Quantity": quantity,
            "Item Cost × Quantity": basic,
            "Order Shipping Amount": shipping,
            "Cart Discount Amount": discount,
            "Order Total Tax Amount": tax,
            **({"Order Total Amount": order_total} if order_total is not None else {}),
        },
    )


def test_dsg_summary_includes_only_completed_and_charges_order_values_once():
    rows = [
        ("DSG", _dsg_summary_row("ORD-1", "Completed", "Book A", 100, 1, 25, 10, 18, 999), "2026-08"),
        ("DSG", _dsg_summary_row("ORD-1", "Completed", "Book B", 200, 2, 25, 20, 18), "2026-08"),
        ("DSG", _dsg_summary_row("ORD-2", "Cancelled", "Book C", 999, 1, 50, 5, 100), "2026-08"),
    ]
    workbook = Workbook()
    workbook.remove(workbook.active)
    _append_dsg_summary_sheet(workbook, rows, "Aug - 26 DSG Sales")
    sheet = workbook["DSG"]

    assert [cell.value for cell in sheet[2]] == [
        "SL No", "Year", "Month", "Order ID", "Type", "Description", "Quantity",
        "Basic Value", "Shipping", "Discount", "Taxable Value",
        "Total Tax", "Total Invoice Value",
    ]
    assert sheet.max_row == 5  # title, header, two completed rows, grand total
    assert [sheet.cell(3, column).value for column in range(1, 14)] == [
        1, 2026, "August", "ORD-1", "Books", "Book A", 1.0, 100.0, 25.0, 10.0, 115.0, 18.0, 133.0,
    ]
    assert [sheet.cell(4, column).value for column in range(1, 14)] == [
        2, 2026, "August", "ORD-1", "Books", "Book B", 2.0, 200.0, 0.0, 20.0, 180.0, 0.0, 180.0,
    ]


def test_dsg_pnl_amount_adds_shipping_once_and_discount_per_product_row():
    first = _dsg_summary_row("ORD-1", "Completed", "Book A", 100, 1, 25, -10, 18)
    second = _dsg_summary_row("ORD-1", "Completed", "Book B", 200, 1, 25, -20, 18)
    charged_orders: set[str] = set()

    assert _dsg_pnl_amount(first, charged_orders) == 115
    assert _dsg_pnl_amount(second, charged_orders) == 180
    assert _dsg_pnl_amount(
        _dsg_summary_row("ORD-2", "Completed", "Book C", 300, 1, 30, -15, 0),
        charged_orders,
    ) == 315


def test_dsg_basic_amount_uses_item_cost_x_quantity_only():
    row = _dsg_summary_row("ORD-1", "Completed", "Book A", 100, 1, 25, -10, 18)
    row.amount = "999"

    assert _dsg_basic_amount(row) == 100


def test_financial_breakdown_uses_channel_specific_columns_and_statuses(monkeypatch):
    def record(**values):
        return SimpleNamespace(upload_id="upload", amount="999", **values)

    rows = {
        DSGDatasetRow: [
            record(order_number="D1", row_data={"Order Status": "Completed", "iteamcostxquantity": 100, "Order Shipping Amount": 10, "Cart Discount Amount": 5, "Taxable Value": 105, "Order Total Tax Amount": 18, "Order Total Amount": 123}),
            record(order_number="D1", row_data={"Order Status": "Completed", "iteamcostxquantity": 50, "Order Shipping Amount": 10, "Cart Discount Amount": 2, "Taxable Value": 58, "Order Total Tax Amount": 18, "Order Total Amount": 123}),
            record(order_number="D2", row_data={"Order Status": "Cancelled", "iteamcostxquantity": 900, "Taxable Value": 900, "Order Total Amount": 900}),
        ],
        SFHDatasetRow: [
            record(row_data={"Invoice No.": "S1", "Earnings Currency": "₹", "Without Tax Total": 200, "Earnings": "$999", "Tax": 36}),
            record(row_data={"Invoice No.": "S2", "Earnings Currency": "INR", "Without Tax Total": 800, "Earnings": "$50.25", "Tax": 99}),
        ],
        AmazonDatasetRow: [
            record(row_data={"amazon-order-id": "A1", "order-status": "Shipped - Delivered to Buyer", "item-price": 300, "shipping-price": 20}),
            record(row_data={"amazon-order-id": "A1", "order-status": "Shipped - Delivered to Buyer", "item-price": 100, "shipping-price": 20}),
            record(row_data={"amazon-order-id": "A2", "order-status": "Shipped", "item-price": 700, "shipping-price": 30}),
        ],
        DirectSalesDatasetRow: [
            record(order_number="I1", row_data={"Status": "Paid", "Without Tax Total": 60, "Discount": 10, "Tax": 18, "Total": 108}),
            record(order_number="I1", row_data={"Status": "Paid", "Without Tax Total": 40, "Discount": 10, "Tax": 18, "Total": 108}),
            record(order_number="I2", row_data={"Status": "Cancelled", "Without Tax Total": 500, "Discount": 0, "Tax": 90, "Total": 590}),
        ],
    }
    monkeypatch.setattr("app.dashboard._cached_rows", lambda model: rows[model])
    monkeypatch.setattr("app.dashboard._cached_upload_dates", lambda: {"upload": datetime(2026, 8, 1)})

    result = _financial_breakdown("monthly", "8", 2026)

    assert result["DSG"] == {"basic_value": 150, "shipping": 10, "discount": 7, "taxable_value": 153, "total_tax": 18, "total_sale": 171}
    assert result["SFH"] == {"basic_value": 250.25, "shipping": 0, "discount": 0, "taxable_value": 250.25, "total_tax": 36, "total_sale": 286.25}
    assert result["Amazon"] == {"basic_value": 400, "shipping": 40, "discount": 0, "taxable_value": 440, "total_tax": 0, "total_sale": 440}
    assert result["Direct Sales"] == {"basic_value": 100, "shipping": 0, "discount": 0, "taxable_value": 100, "total_tax": 18, "total_sale": 108}


def test_dsg_summary_invoice_value_is_taxable_value_plus_tax():
    first = _dsg_summary_row("ORD-1", "Completed", "Book A", 100, 1, 25, 0, 18)
    second = _dsg_summary_row("ORD-1", "Completed", "Book B", 200, 1, 25, 0, 18)
    first.row_data["Order Total Amount"] = 344
    second.row_data["Order Total Amount"] = 344
    workbook = Workbook()
    workbook.remove(workbook.active)

    _append_dsg_summary_sheet(workbook, [("DSG", first, "2026-08"), ("DSG", second, "2026-08")], "Aug - 26 DSG Sales")

    sheet = workbook["DSG"]
    assert sheet.cell(3, 13).value == 143
    assert sheet.cell(4, 13).value == 200


def _sfh_summary_row(invoice, currency, course, without_tax, earnings, tax):
    return SimpleNamespace(
        course=course,
        product_name=course,
        category="Web Version",
        row_data={
            "Invoice No.": invoice,
            "Currency": currency,
            "Course": course,
            "Without Tax Total": without_tax,
            "Earnings": earnings,
            "Tax": tax,
        },
    )


def test_sfh_summary_deduplicates_invoice_and_applies_currency_rules():
    rows = [
        ("SFH", _sfh_summary_row("INV-1", "₹", "Course A", 100, 150, 18), "2026-08"),
        ("SFH", _sfh_summary_row("INV-1", "₹", "Duplicate", 100, 150, 18), "2026-08"),
        ("SFH", _sfh_summary_row("INV-2", "INR", "Course B", 200, 250, 36), "2026-08"),
        ("SFH", _sfh_summary_row("INV-3", "USD", "Course C", 300, 275, 54), "2026-08"),
    ]
    workbook = Workbook()
    workbook.remove(workbook.active)
    _append_sfh_summary_sheet(workbook, rows, "Aug - 26 SFH Sales")
    sheet = workbook["SFH"]

    assert sheet.max_row == 6  # title, header, three unique invoices, grand total
    assert [sheet.cell(3, column).value for column in range(1, 14)] == [
        1, 2026, "August", "INV-1", "Web Version", "Course A", 1, 100.0, 0, 0, 100.0, 18.0, 118.0,
    ]
    assert [sheet.cell(4, column).value for column in range(1, 14)] == [
        2, 2026, "August", "INV-2", "Web Version", "Course B", 1, 250.0, 0, 0, 250.0, 0.0, 250.0,
    ]
    assert [sheet.cell(5, column).value for column in range(1, 14)] == [
        3, 2026, "August", "INV-3", "Web Version", "Course C", 1, 275.0, 0, 0, 275.0, 0.0, 275.0,
    ]


def _amazon_summary_row(order, status, product, quantity, item_price, shipping):
    return SimpleNamespace(
        product_name=product,
        category="Books",
        amount=str(item_price),
        row_data={
            "amazon-order-id": order,
            "order-status": status,
            "product-name": product,
            "quantity": quantity,
            "item-price": item_price,
            "shipping-price": shipping,
        },
    )


def test_amazon_summary_excludes_cancelled_and_calculates_without_tax_or_discount():
    rows = [
        ("Amazon", _amazon_summary_row("AMZ-1", "Shipped - Delivered to Buyer", "Book A", 2, 500, 40), "2026-08"),
        ("Amazon", _amazon_summary_row("AMZ-2", "Cancelled", "Book B", 1, 999, 50), "2026-08"),
        ("Amazon", _amazon_summary_row("AMZ-3", "Shipped - Delivered to Buyer", "Book C", 3, 300, 0), "2026-08"),
    ]
    workbook = Workbook()
    workbook.remove(workbook.active)
    _append_amazon_summary_sheet(workbook, rows, "Aug - 26 Amazon Sales")
    sheet = workbook["Amazon"]

    assert sheet.max_row == 5  # title, header, two valid rows, grand total
    assert [sheet.cell(3, column).value for column in range(1, 14)] == [
        1, 2026, "August", "AMZ-1", "Books", "Book A", 2.0, 500.0, 40.0, 0, 540.0, 0, 540.0,
    ]
    assert [sheet.cell(4, column).value for column in range(1, 14)] == [
        2, 2026, "August", "AMZ-3", "Books", "Book C", 3.0, 300.0, 0.0, 0, 300.0, 0, 300.0,
    ]


def test_amazon_channel_performance_uses_delivered_unique_orders_and_item_price():
    rows = [
        ("Amazon", _amazon_summary_row("AMZ-1", "Shipped - Delivered to Buyer", "A", 1, 100, 0), "2026-04"),
        ("Amazon", _amazon_summary_row("AMZ-1", "Shipped - Delivered to Buyer", "B", 1, 50, 0), "2026-04"),
        ("Amazon", _amazon_summary_row("AMZ-2", " shipped - delivered to buyer ", "C", 1, 75, 0), "2026-04"),
        ("Amazon", _amazon_summary_row("AMZ-3", "Cancelled", "D", 1, 999, 0), "2026-04"),
        ("Amazon", _amazon_summary_row("AMZ-4", "Shipped - Delivered to Buyer", "E", 1, 125, 0), "2026-05"),
        ("DSG", _dsg_summary_row("DSG-1", "Completed", "Book", 500, 1, 0, 0, 0), "2026-05"),
    ]

    sheet = _channel_performance_workbook(rows)["Amazon"]

    assert [sheet.cell(2, column).value for column in (3, 4, 5, 6)] == [2, "-", "-", 225]
    assert [sheet.cell(3, column).value for column in (3, 4, 5, 6)] == [1, -1, -0.5, 125]


def test_amazon_channel_performance_handles_zero_previous_orders():
    rows = [
        ("Amazon", _amazon_summary_row("", "Shipped - Delivered to Buyer", "A", 1, 100, 0), "2026-04"),
        ("Amazon", _amazon_summary_row("AMZ-1", "Shipped - Delivered to Buyer", "B", 1, 50, 0), "2026-05"),
    ]

    sheet = _channel_performance_workbook(rows)["Amazon"]

    assert sheet.cell(2, 3).value == 0
    assert sheet.cell(3, 4).value == 1
    assert sheet.cell(3, 5).value == "-"


def test_unique_order_count_uses_channel_rules_and_amazon_exclusions():
    completed = _dsg_summary_row("SAME-ID", "Completed", "A", 1, 1, 0, 0, 0)
    cancelled = _dsg_summary_row("DSG-X", "Cancelled", "B", 1, 1, 0, 0, 0)
    sfh = _sfh_summary_row("SAME-ID", "INR", "Course", 1, 1, 0)
    returning = _amazon_summary_row("AMZ-X", "Shipped - Returning to Seller", "A", 1, 1, 0)
    shipped = _amazon_summary_row("AMZ-OK", "Shipped - Delivered to Buyer", "B", 1, 1, 0)
    rows = [
        ("DSG", completed, "2026-08"), ("DSG", completed, "2026-08"),
        ("DSG", cancelled, "2026-08"), ("SFH", sfh, "2026-08"),
        ("SFH", sfh, "2026-08"), ("Amazon", returning, "2026-08"),
        ("Amazon", shipped, "2026-08"),
    ]

    assert _unique_order_count(rows) == 3
    assert _amazon_is_cancelled(returning) is True


def test_amazon_zero_item_price_is_excluded_from_orders_and_reports():
    zero_price = _amazon_summary_row("AMZ-ZERO", "Shipped - Delivered to Buyer", "Free item", 1, 0, 50)
    na_price = _amazon_summary_row("AMZ-NA", "Shipped - Delivered to Buyer", "Unavailable item", 1, "N/A", 0)
    wrong_status = _amazon_summary_row("AMZ-SHIPPED", "Shipped", "Undelivered item", 1, 200, 0)
    paid = _amazon_summary_row("AMZ-PAID", "Shipped - Delivered to Buyer", "Paid item", 1, 125, 0)
    rows = [
        ("Amazon", zero_price, "2026-08"),
        ("Amazon", na_price, "2026-08"),
        ("Amazon", wrong_status, "2026-08"),
        ("Amazon", paid, "2026-08"),
    ]

    assert _amazon_is_excluded(zero_price) is True
    assert _amazon_is_excluded(na_price) is True
    assert _amazon_is_excluded(wrong_status) is True
    assert _unique_order_count(rows) == 1

    workbook = Workbook()
    workbook.remove(workbook.active)
    _append_amazon_summary_sheet(workbook, rows, "Aug - 26 Amazon Sales")
    sheet = workbook["Amazon"]
    assert sheet.max_row == 4
    assert sheet.cell(3, 4).value == "AMZ-PAID"
    assert sheet.cell(3, 13).value == 125


def test_direct_sales_summary_maps_fields_and_allocates_invoice_values_without_duplication():
    rows = [
        ("Direct Sales", SimpleNamespace(
            order_number="INV-1", category="Books", product_name="Book A", amount="60",
            row_data={
                "Invoice Number": "INV-1", "Without Tax Total": 60,
                "Original Without Tax Total": 100, "Tax": 18,
                "Category": "Books", "Item Details": "Book A", "Category Quantity": 2,
            },
        ), "2026-08"),
        ("Direct Sales", SimpleNamespace(
            order_number="INV-1", category="Audio Device", product_name="Player", amount="40",
            row_data={
                "Invoice Number": "INV-1", "Without Tax Total": 40,
                "Original Without Tax Total": 100, "Tax": 18,
                "Category": "Audio Device", "Item Details": "Player", "Category Quantity": 1,
            },
        ), "2026-08"),
        ("Direct Sales", SimpleNamespace(
            order_number="INV-2", category="Pen Drive", product_name="USB Course", amount="50",
            row_data={
                "Invoice Number": "INV-2", "Without Tax Total": 50, "Tax": 9,
                "Category": "Pen Drive", "Item Details": "USB Course", "Category Quantity": 3,
            },
        ), "2026-08"),
    ]
    workbook = Workbook()
    workbook.remove(workbook.active)
    _append_direct_sales_summary_sheet(workbook, rows, "Aug - 26 Direct Sales")
    sheet = workbook["Direct Sales"]

    assert [cell.value for cell in sheet[2]] == [
        "SL No", "Year", "Month", "Order ID", "Type", "Description", "Quantity",
        "Taxable Value", "Total Tax", "Total Invoice Value",
    ]
    assert [sheet.cell(3, column).value for column in range(1, 11)] == [
        1, 2026, "August", "INV-1", "Books", "Book A", 2.0, 60.0, 10.8, 70.8,
    ]
    assert [sheet.cell(4, column).value for column in range(1, 11)] == [
        2, 2026, "August", "INV-1", "Audio Device", "Player", 1.0, 40.0, 7.2, 47.2,
    ]
    assert [sheet.cell(5, column).value for column in range(1, 11)] == [
        3, 2026, "August", "INV-2", "Pen Drive", "USB Course", 3.0, 50.0, 9.0, 59.0,
    ]
    assert sum(sheet.cell(row, 8).value for row in (3, 4)) == 100.0
    assert sum(sheet.cell(row, 9).value for row in (3, 4)) == 18.0


class _RowsQuery:
    def __init__(self, rows):
        self.rows = rows

    def all(self):
        return self.rows


class _OverviewDatabase:
    def __init__(self, direct_rows, uploads):
        self.direct_rows = direct_rows
        self.uploads = uploads

    def query(self, model):
        return _RowsQuery(self.uploads if model is UploadHistory else self.direct_rows)


def _direct_row(row_id, order, product, amount, quantity, notes=""):
    return SimpleNamespace(
        id=row_id,
        upload_id="upload-1",
        order_number=order,
        product_name=product,
        category="Books",
        amount=str(amount),
        row_data={
            "Issue Date": "2026-04-10",
            "Without Tax Total": amount,
            "Mapped Quantity": quantity,
            "Category Quantity": quantity,
            "Private Notes": notes,
        },
    )


def test_direct_sales_amount_excludes_cancelled_invoice_status():
    active = _direct_row(1, "A", "Active Book", 100, 1)
    cancelled = _direct_row(2, "B", "Cancelled Book", 250, 1)
    cancelled.row_data["Status"] = "  Cancelled "

    assert _direct_amount(active) == 100
    assert _direct_amount(cancelled) == 0


def test_direct_sales_overview_reuses_dashboard_mapping_and_reconciles():
    database = _OverviewDatabase(
        [
            _direct_row(1, "A", "Bulk Book", 100.4, 11),
            _direct_row(2, "B", "Retail Book", 50.4, 2),
            _direct_row(3, "C", "Stall Book", 25.4, 1, "Stall counter"),
        ],
        [SimpleNamespace(upload_id="upload-1", uploaded_at=None)],
    )
    result = _build_direct_sales_overview(
        database,
        years={2026},
        months={4},
        types={"Bulk", "In Office", "Stall"},
        products=set(),
    )

    assert result["validated"] is True
    assert result["totals"] == {
        "In Office": 50, "Stall": 25, "Bulk": 101,
        "Call": 0, "Retail": 0, "Language Lab": 0, "Course Promotion": 0,
    }
    assert result["overall_total"] == 176
    assert {row["type"] for row in result["rows"]} == {"Bulk", "In Office", "Stall"}


def test_direct_sales_overview_reconciles_type_totals_to_overall_total():
    rows = [
        _direct_row(1, "B-1", "Bulk item", 4186.9, 11),
        _direct_row(2, "R-1", "Retail item", 69699.3, 1),
        _direct_row(3, "S-1", "Stall item", 38175.37, 1, "STALL"),
    ]
    database = _OverviewDatabase(
        rows,
        [SimpleNamespace(upload_id="upload-1", uploaded_at=None)],
    )

    result = _build_direct_sales_overview(
        database,
        years={2026}, months={4}, types={"Bulk", "In Office", "Stall"}, products=set(),
    )

    assert result["overall_total"] == 112062
    assert sum(result["totals"].values()) == 112062


def test_direct_sales_overview_separates_language_lab_from_retail():
    database = _OverviewDatabase(
        [_direct_row(1, "L-1", "Language Lab", 48813.57, 1, "Language Lab")],
        [SimpleNamespace(upload_id="upload-1", uploaded_at=None)],
    )

    result = _build_direct_sales_overview(
        database,
        years={2026}, months={4}, types={"Language Lab"}, products=set(),
    )

    assert result["totals"]["Language Lab"] == 48814
    assert result["totals"]["Retail"] == 0
    assert result["overall_total"] == 48814


def test_direct_sales_overview_includes_course_promotion_in_priority_order():
    database = _OverviewDatabase(
        [
            _direct_row(1, "I-1", "Office Book", 100, 1),
            _direct_row(2, "C-1", "S101 Course", 250, 1, "course promotion"),
            _direct_row(3, "S-1", "Stall Book", 300, 1, "stall"),
        ],
        [SimpleNamespace(upload_id="upload-1", uploaded_at=None)],
    )

    result = _build_direct_sales_overview(
        database,
        years={2026}, months={4},
        types={"Stall", "Course Promotion", "In Office"}, products=set(),
    )

    assert [row["type"] for row in result["rows"]] == [
        "Stall", "Course Promotion", "In Office",
    ]
    assert result["totals"]["Course Promotion"] == 250
    assert result["overall_total"] == 650
    assert result["detail_rows"] == [
        {"year": 2026, "month": "Apr", "invoice": "s-1", "type": "Stall", "product": "Stall Book", "quantity": 1.0, "without_tax_total": 300.0},
        {"year": 2026, "month": "Apr", "invoice": "c-1", "type": "Course Promotion", "product": "S101 Course", "quantity": 1.0, "without_tax_total": 250.0},
        {"year": 2026, "month": "Apr", "invoice": "i-1", "type": "In Office", "product": "Office Book", "quantity": 1.0, "without_tax_total": 100.0},
    ]


def test_direct_sales_overview_exposes_call_retail_and_in_office_separately():
    database = _OverviewDatabase(
        [
            _direct_row(1, "I-1", "Office Book", 100, 1),
            _direct_row(2, "C-1", "Call Book", 200, 1, "phone call"),
            _direct_row(3, "R-1", "Retail Book", 300, 1, "vendant"),
        ],
        [SimpleNamespace(upload_id="upload-1", uploaded_at=None)],
    )

    result = _build_direct_sales_overview(
        database,
        years={2026}, months={4},
        types={"In Office", "Call", "Retail"}, products=set(),
    )

    assert result["totals"] == {
        "In Office": 100, "Stall": 0, "Bulk": 0,
        "Call": 200, "Retail": 300, "Language Lab": 0, "Course Promotion": 0,
    }


def test_channel_performance_keeps_call_and_retail_separate():
    assert _direct_sales_channel(_direct_row(1, "C-1", "Call", 100, 1, "phone call")) == "Call"
    assert _direct_sales_channel(_direct_row(2, "R-1", "Retail", 100, 1, "vendant")) == "Retail"
    assert _direct_sales_channel(_direct_row(5, "R-2", "Retail Bulk Qty", 100, 25, "vendant")) == "Retail"
    assert _direct_sales_channel(_direct_row(3, "S-1", "Stall", 100, 1, "stall")) == "Stall"
    assert _direct_sales_channel(_direct_row(4, "B-1", "Bulk", 100, 11)) == "Bulk"


def test_direct_sales_overview_quantity_over_ten_is_always_bulk():
    row = _direct_row(1, "B-1", "Course Book", 100, 11, "stall course call vendant")

    assert _direct_overview_type(row) == "Bulk"
    assert _direct_sales_channel(row) == "Stall"


def test_direct_sales_overview_rejects_duplicate_source_rows():
    row = _direct_row(1, "A", "Book", 100, 1)
    database = _OverviewDatabase(
        [row, row],
        [SimpleNamespace(upload_id="upload-1", uploaded_at=None)],
    )
    with pytest.raises(Exception, match="Duplicate Direct Sales database row"):
        _build_direct_sales_overview(
            database,
            years={2026}, months={4}, types={"Retail"}, products=set(),
        )
